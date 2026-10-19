import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from tallycraft import AccountType, EntryLine, Ledger
from tallycraft.ledger import AccountingError


class ExcelInterchangeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = Ledger(self.root / "source.sqlite3")
        self._add_chart(self.source)

    def tearDown(self):
        self.source.close()
        self.temp.cleanup()

    @staticmethod
    def _add_chart(ledger):
        for code, name, kind in (
            ("1000", "Cash", AccountType.ASSET),
            ("1200", "Inventory", AccountType.ASSET),
            ("2000", "Accounts payable", AccountType.LIABILITY),
            ("3000", "Equity", AccountType.EQUITY),
            ("4000", "Sales", AccountType.INCOME),
            ("5000", "Cost of goods sold", AccountType.EXPENSE),
        ):
            ledger.create_account(code, name, kind)

    def _add_inventory_transactions(self):
        self.source.create_inventory_item(
            "BOOK-1", "Accounting notebook", "pcs", "1200", "5000", "2"
        )
        self.source.post_entry(
            "2026-03-01",
            "Initial funds",
            [EntryLine("1000", debit="100.00"), EntryLine("3000", credit="100.00")],
            "OPEN-1",
        )
        self.source.receive_inventory(
            "BOOK-1", "2026-03-02", "10", "3.00", "2000", "Supplier receipt", "GRN-1"
        )
        self.source.receive_inventory(
            "BOOK-1", "2026-03-03", "10", "5.00", "2000", "Supplier receipt", "GRN-2"
        )
        self.source.issue_inventory("BOOK-1", "2026-03-04", "5", "Stock issued", "ISS-1")

    def test_xlsx_round_trip_preserves_accounts_journals_and_inventory(self):
        self._add_inventory_transactions()
        path = self.root / "company.xlsx"
        self.source.export_excel(path)
        target = Ledger(self.root / "target.sqlite3")
        try:
            self._add_chart(target)
            imported_vouchers = target.import_excel(path)
            self.assertEqual(imported_vouchers, 4)
            self.assertEqual(target.list_accounts(), self.source.list_accounts())
            self.assertEqual(target.trial_balance("2026-03-31"), self.source.trial_balance("2026-03-31"))
            self.assertEqual(target.list_inventory_items(), self.source.list_inventory_items())
            source_movements = self.source.inventory_history()
            imported_movements = target.inventory_history()
            self.assertEqual(
                [(item.movement_type, item.quantity, item.value) for item in imported_movements],
                [(item.movement_type, item.quantity, item.value) for item in source_movements],
            )
            target.verify_integrity()
        finally:
            target.close()

    def test_invalid_workbook_rolls_back_accounts_and_vouchers(self):
        self.source.create_account("9000", "Imported account", AccountType.ASSET)
        self.source.post_entry(
            "2026-03-01",
            "Initial funds",
            [EntryLine("1000", debit="100.00"), EntryLine("3000", credit="100.00")],
        )
        path = self.root / "tampered.xlsx"
        self.source.export_excel(path)
        workbook = load_workbook(path)
        workbook["Voucher Lines"]["D2"] = Decimal("200")
        workbook.save(path)
        workbook.close()

        target = Ledger(self.root / "target.sqlite3")
        try:
            self._add_chart(target)
            with self.assertRaisesRegex(AccountingError, "not balanced"):
                target.import_excel(path)
            self.assertEqual(
                [account.code for account in target.list_accounts()],
                ["1000", "1200", "2000", "3000", "4000", "5000"],
            )
            self.assertEqual(target.journal_history(), ())
        finally:
            target.close()

    def test_import_refuses_nonempty_book_to_avoid_duplicate_postings(self):
        self.source.post_entry(
            "2026-03-01",
            "Initial funds",
            [EntryLine("1000", debit="100.00"), EntryLine("3000", credit="100.00")],
        )
        path = self.root / "company.xlsx"
        self.source.export_excel(path)
        target = Ledger(self.root / "target.sqlite3")
        try:
            self._add_chart(target)
            target.post_entry(
                "2026-03-02",
                "Already posted",
                [EntryLine("1000", debit="1"), EntryLine("3000", credit="1")],
            )
            with self.assertRaisesRegex(AccountingError, "no journal entries"):
                target.import_excel(path)
            self.assertEqual(len(target.journal_history()), 1)
        finally:
            target.close()

    def test_invalid_zip_is_reported_as_workbook_error(self):
        path = self.root / "not-an-excel-file.xlsx"
        path.write_text("not an xlsx archive", encoding="utf-8")
        with self.assertRaisesRegex(AccountingError, "Could not read Excel workbook"):
            self.source.import_excel(path)

    def test_export_writes_formula_shaped_account_names_as_literal_text(self):
        self.source.create_account(
            "9000", "=HYPERLINK(\"https://example.invalid\")", AccountType.ASSET
        )
        path = self.root / "formula-safe.xlsx"
        self.source.export_excel(path)
        workbook = load_workbook(path, data_only=False)
        try:
            sheet = workbook["Accounts"]
            matching_row = next(
                row for row in range(2, sheet.max_row + 1)
                if sheet.cell(row, 1).value == "9000"
            )
            cell = sheet.cell(matching_row, 2)
            self.assertEqual(cell.value, '=HYPERLINK("https://example.invalid")')
            self.assertEqual(cell.data_type, "s")
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
