"""Excel workbook interchange for a complete TallyCraft book."""

from __future__ import annotations

import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING

from .ledger import AccountType, AccountingError, Ledger

if TYPE_CHECKING:
    from openpyxl import Workbook

_SHEETS = {
    "Accounts": ("Code", "Name", "Type"),
    "Vouchers": ("Source ID", "Date", "Memo", "Reference"),
    "Voucher Lines": ("Source ID", "Line", "Account Code", "Debit", "Credit"),
    "Inventory Items": (
        "SKU",
        "Name",
        "Unit",
        "Inventory Account",
        "COGS Account",
        "Reorder Level",
        "Quantity On Hand",
        "Carrying Value",
    ),
    "Stock Movements": (
        "Source ID",
        "Date",
        "SKU",
        "Type",
        "Quantity",
        "Value",
        "Reference",
        "Memo",
    ),
    "Trial Balance": ("Code", "Account", "Type", "Debit", "Credit"),
}


def export_workbook(ledger: Ledger, path: str | Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as error:
        raise AccountingError(
            "Excel support requires openpyxl. Install the application requirements."
        ) from error

    ledger.verify_integrity()
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, headers in _SHEETS.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(headers)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1B806D")

    accounts_sheet = workbook["Accounts"]
    for account in ledger.list_accounts():
        _append_row(accounts_sheet, (account.code, account.name, account.type))

    vouchers_sheet = workbook["Vouchers"]
    vouchers = {entry.id: entry for entry in ledger.journal_history(limit=1_000_000)}
    for entry in sorted(vouchers.values(), key=lambda item: item.id):
        _append_row(
            vouchers_sheet, (entry.id, entry.entry_date, entry.memo, entry.reference)
        )

    lines_sheet = workbook["Voucher Lines"]
    for line in ledger.journal_line_details():
        _append_row(
            lines_sheet,
            (line.entry_id, line.line_number, line.account_code, line.debit, line.credit),
        )

    inventory_sheet = workbook["Inventory Items"]
    for item in ledger.list_inventory_items():
        _append_row(
            inventory_sheet,
            (
                item.sku,
                item.name,
                item.unit,
                item.inventory_account_code,
                item.cogs_account_code,
                item.reorder_level,
                item.quantity_on_hand,
                item.carrying_value,
            ),
        )

    movements_sheet = workbook["Stock Movements"]
    for movement in sorted(
        ledger.inventory_history(1_000_000), key=lambda item: item.journal_entry_id
    ):
        _append_row(
            movements_sheet,
            (
                movement.journal_entry_id,
                movement.entry_date,
                movement.sku,
                movement.movement_type,
                movement.quantity,
                movement.value,
                movement.reference,
                movement.memo,
            ),
        )

    trial_balance = ledger.trial_balance(date.today())
    balance_sheet = workbook["Trial Balance"]
    for account, debit, credit in trial_balance.rows:
        _append_row(balance_sheet, (account.code, account.name, account.type, debit, credit))
    _append_row(
        balance_sheet,
        ("", "", "Total", trial_balance.debit_total, trial_balance.credit_total),
    )

    for sheet in workbook.worksheets:
        for column in sheet.columns:
            letter = column[0].column_letter
            max_width = max(
                (len(str(cell.value)) for cell in column if cell.value is not None),
                default=10,
            )
            sheet.column_dimensions[letter].width = min(max(max_width + 2, 12), 42)

    try:
        workbook.save(path)
    except (OSError, ValueError) as error:
        raise AccountingError(f"Could not export Excel workbook: {error}") from error
    finally:
        workbook.close()


def import_workbook(ledger: Ledger, path: str | Path) -> int:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise AccountingError(
            "Excel support requires openpyxl. Install the application requirements."
        ) from error

    try:
        from openpyxl.utils.exceptions import InvalidFileException

        workbook = load_workbook(path, read_only=True, data_only=False)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, InvalidFileException) as error:
        raise AccountingError(f"Could not read Excel workbook: {error}") from error

    try:
        missing_sheets = set(_SHEETS) - set(workbook.sheetnames)
        if missing_sheets:
            raise AccountingError(
                "Workbook is missing required sheet(s): "
                + ", ".join(sorted(missing_sheets))
            )
        tables = {
            name: _read_sheet(workbook[name], headers)
            for name, headers in _SHEETS.items()
        }
    finally:
        workbook.close()

    return _import_tables(ledger, tables)


def _append_row(sheet: object, values: tuple[object, ...]) -> None:
    row_number = sheet.max_row + 1
    for column_number, value in enumerate(values, start=1):
        cell = sheet.cell(row=row_number, column=column_number, value=value)
        if isinstance(value, str):
            cell.data_type = "s"


def _read_sheet(sheet: object, expected_headers: tuple[str, ...]) -> list[tuple[object, ...]]:
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = next(rows)
    except StopIteration as error:
        raise AccountingError(f"Sheet {sheet.title} is empty.") from error
    if tuple(headers[: len(expected_headers)]) != expected_headers:
        raise AccountingError(
            f"Sheet {sheet.title} must start with columns: "
            + ", ".join(expected_headers)
        )
    result = []
    for row_number, row in enumerate(rows, start=2):
        values = tuple(row[: len(expected_headers)])
        if all(value is None for value in values):
            continue
        if len(values) != len(expected_headers):
            raise AccountingError(f"Sheet {sheet.title}, row {row_number}: invalid row.")
        result.append(values)
    return result


def _ensure_empty_book(ledger: Ledger) -> None:
    if ledger._db.execute("SELECT EXISTS(SELECT 1 FROM journal_entries)").fetchone()[0]:
        raise AccountingError(
            "Import is only allowed into a book with no journal entries, "
            "to prevent duplicate postings."
        )
    if ledger._db.execute("SELECT EXISTS(SELECT 1 FROM inventory_items)").fetchone()[0]:
        raise AccountingError(
            "Import is only allowed into a book with no inventory items."
        )


def _import_tables(ledger: Ledger, tables: dict[str, list[tuple[object, ...]]]) -> int:
    vouchers = _parse_vouchers(tables["Vouchers"])
    lines_by_source = _parse_lines(tables["Voucher Lines"], set(vouchers))
    if set(lines_by_source) != set(vouchers):
        missing = sorted(set(vouchers) - set(lines_by_source))
        raise AccountingError(f"Voucher(s) have no lines: {', '.join(map(str, missing))}")
    if tables["Stock Movements"] and not tables["Inventory Items"]:
        raise AccountingError("Stock movements require matching inventory item records.")

    imported_vouchers = 0
    with ledger._db:
        _ensure_empty_book(ledger)
        account_ids = {
            row["code"]: row["id"]
            for row in ledger._db.execute("SELECT id, code FROM accounts")
        }
        account_types = {
            row["code"]: row["type"]
            for row in ledger._db.execute("SELECT code, type FROM accounts")
        }
        _import_accounts(ledger, tables["Accounts"], account_ids, account_types)

        source_entries: dict[int, int] = {}
        posted_lines: dict[int, list[tuple[str, int, int]]] = {}
        for source_id, voucher in vouchers.items():
            normalized = lines_by_source[source_id]
            _validate_balanced_lines(source_id, normalized)
            new_id = ledger._post_entry_transaction(
                voucher["date"], voucher["memo"], voucher["reference"], normalized
            )
            source_entries[source_id] = new_id
            posted_lines[source_id] = normalized
            imported_vouchers += 1

        _import_inventory(
            ledger,
            tables["Inventory Items"],
            tables["Stock Movements"],
            source_entries,
            posted_lines,
        )
        ledger.verify_integrity()
    return imported_vouchers


def _import_accounts(
    ledger: Ledger,
    rows: list[tuple[object, ...]],
    account_ids: dict[str, int],
    account_types: dict[str, str],
) -> None:
    seen: set[str] = set()
    for row_number, (raw_code, raw_name, raw_type) in enumerate(rows, start=2):
        code = _text(raw_code, "Accounts", row_number, "Code")
        name = _text(raw_name, "Accounts", row_number, "Name")
        account_type = _text(raw_type, "Accounts", row_number, "Type").lower()
        if code in seen:
            raise AccountingError(f"Accounts row {row_number}: duplicate code {code}.")
        seen.add(code)
        if account_type not in AccountType.VALUES:
            raise AccountingError(
                f"Accounts row {row_number}: unsupported account type {account_type}."
            )
        if code in account_ids:
            existing = ledger._db.execute(
                "SELECT name, type FROM accounts WHERE code = ?", (code,)
            ).fetchone()
            if (existing["name"], existing["type"]) != (name, account_type):
                raise AccountingError(
                    f"Accounts row {row_number}: code {code} conflicts with the existing chart."
                )
            continue
        cursor = ledger._db.execute(
            "INSERT INTO accounts(code, name, type) VALUES (?, ?, ?)",
            (code, name, account_type),
        )
        account_ids[code] = cursor.lastrowid
        account_types[code] = account_type


def _parse_vouchers(rows: list[tuple[object, ...]]) -> dict[int, dict[str, object]]:
    vouchers = {}
    for row_number, (raw_id, raw_date, raw_memo, raw_reference) in enumerate(rows, start=2):
        source_id = _source_id(raw_id, "Vouchers", row_number)
        if source_id in vouchers:
            raise AccountingError(f"Vouchers row {row_number}: duplicate source ID {source_id}.")
        memo = _text(raw_memo, "Vouchers", row_number, "Memo")
        reference = None if _is_blank(raw_reference) else str(raw_reference).strip()
        vouchers[source_id] = {
            "date": _date(raw_date, "Vouchers", row_number),
            "memo": memo,
            "reference": reference,
        }
    return vouchers


def _parse_lines(
    rows: list[tuple[object, ...]], source_ids: set[int]
) -> dict[int, list[tuple[str, int, int]]]:
    result: dict[int, list[tuple[str, int, int]]] = {}
    seen_lines: set[tuple[int, int]] = set()
    for row_number, (raw_id, raw_number, raw_code, raw_debit, raw_credit) in enumerate(
        rows, start=2
    ):
        source_id = _source_id(raw_id, "Voucher Lines", row_number)
        if source_id not in source_ids:
            raise AccountingError(
                f"Voucher Lines row {row_number}: unknown source ID {source_id}."
            )
        line_number = _source_id(raw_number, "Voucher Lines", row_number, "Line")
        if (source_id, line_number) in seen_lines:
            raise AccountingError(
                f"Voucher Lines row {row_number}: duplicate line number {line_number}."
            )
        seen_lines.add((source_id, line_number))
        code = _text(raw_code, "Voucher Lines", row_number, "Account Code")
        debit = ledger_amount(raw_debit, "Voucher Lines", row_number, "Debit")
        credit = ledger_amount(raw_credit, "Voucher Lines", row_number, "Credit")
        result.setdefault(source_id, []).append((code, debit, credit))
    return result


def ledger_amount(value: object, sheet: str, row: int, column: str) -> int:
    try:
        if isinstance(value, bool) or value is None:
            raise ValueError
        return Ledger._minor(Decimal(str(value)))
    except (AccountingError, ArithmeticError, ValueError, TypeError) as error:
        raise AccountingError(f"{sheet} row {row}: invalid {column} amount.") from error


def _validate_balanced_lines(source_id: int, lines: list[tuple[str, int, int]]) -> None:
    if len(lines) < 2:
        raise AccountingError(f"Voucher {source_id} must have at least two lines.")
    debit_sum = credit_sum = 0
    for code, debit, credit in lines:
        if (debit > 0) == (credit > 0):
            raise AccountingError(
                f"Voucher {source_id} must have exactly one side per journal line."
            )
        debit_sum += debit
        credit_sum += credit
    if debit_sum != credit_sum:
        raise AccountingError(f"Voucher {source_id} is not balanced.")


def _import_inventory(
    ledger: Ledger,
    item_rows: list[tuple[object, ...]],
    movement_rows: list[tuple[object, ...]],
    source_entries: dict[int, int],
    posted_lines: dict[int, list[tuple[str, int, int]]],
) -> None:
    expected: dict[str, tuple[int, int]] = {}
    item_accounts: dict[str, tuple[str, str]] = {}
    item_ids: dict[str, int] = {}
    account_types = {
        row["code"]: row["type"]
        for row in ledger._db.execute("SELECT code, type FROM accounts")
    }
    seen: set[str] = set()
    for row_number, row in enumerate(item_rows, start=2):
        (
            raw_sku, raw_name, raw_unit, raw_inventory_account, raw_cogs_account,
            raw_reorder, raw_quantity, raw_value,
        ) = row
        sku = _text(raw_sku, "Inventory Items", row_number, "SKU")
        name = _text(raw_name, "Inventory Items", row_number, "Name")
        unit = _text(raw_unit, "Inventory Items", row_number, "Unit")
        inventory_code = _text(raw_inventory_account, "Inventory Items", row_number, "Inventory Account")
        cogs_code = _text(raw_cogs_account, "Inventory Items", row_number, "COGS Account")
        if sku in seen:
            raise AccountingError(f"Inventory Items row {row_number}: duplicate SKU {sku}.")
        seen.add(sku)
        if account_types.get(inventory_code) != AccountType.ASSET:
            raise AccountingError(
                f"Inventory Items row {row_number}: inventory account must be an asset."
            )
        if account_types.get(cogs_code) != AccountType.EXPENSE:
            raise AccountingError(
                f"Inventory Items row {row_number}: COGS account must be an expense."
            )
        reorder = Ledger._quantity_milli(_quantity(raw_reorder, row_number, "Reorder Level"), allow_zero=True)
        opening_quantity = Ledger._quantity_milli(_quantity(raw_quantity, row_number, "Quantity On Hand"), allow_zero=True)
        opening_value = ledger_amount(raw_value, "Inventory Items", row_number, "Carrying Value")
        cursor = ledger._db.execute(
            """INSERT INTO inventory_items
               (sku, name, unit, inventory_account_id, cogs_account_id, reorder_milli)
               VALUES (?, ?, ?, ?, ?, ?)""",
            (
                sku, name, unit,
                _account_id(ledger, inventory_code, row_number, "Inventory Account"),
                _account_id(ledger, cogs_code, row_number, "COGS Account"),
                reorder,
            ),
        )
        item_ids[sku] = cursor.lastrowid
        item_accounts[sku] = (inventory_code, cogs_code)
        expected[sku] = (opening_quantity, opening_value)

    movement_seen: set[int] = set()
    for row_number, row in enumerate(movement_rows, start=2):
        (
            raw_source_id, raw_date, raw_sku, raw_type, raw_quantity, raw_value,
            raw_reference, raw_memo,
        ) = row
        source_id = _source_id(raw_source_id, "Stock Movements", row_number)
        if source_id not in source_entries:
            raise AccountingError(
                f"Stock Movements row {row_number}: no matching voucher {source_id}."
            )
        if source_id in movement_seen:
            raise AccountingError(
                f"Stock Movements row {row_number}: a voucher can only represent one stock movement."
            )
        movement_seen.add(source_id)
        sku = _text(raw_sku, "Stock Movements", row_number, "SKU")
        if sku not in item_ids:
            raise AccountingError(f"Stock Movements row {row_number}: unknown SKU {sku}.")
        movement_type = _text(raw_type, "Stock Movements", row_number, "Type").lower()
        if movement_type not in ("receipt", "issue"):
            raise AccountingError(f"Stock Movements row {row_number}: invalid movement type.")
        quantity = Ledger._quantity_milli(_quantity(raw_quantity, row_number, "Quantity"))
        value = ledger_amount(raw_value, "Stock Movements", row_number, "Value")
        if value < 1:
            raise AccountingError(f"Stock Movements row {row_number}: value must be positive.")
        movement_date = _date(raw_date, "Stock Movements", row_number)
        voucher = ledger._db.execute(
            "SELECT entry_date FROM journal_entries WHERE id = ?",
            (source_entries[source_id],),
        ).fetchone()
        if voucher["entry_date"] != movement_date.isoformat():
            raise AccountingError(
                f"Stock Movements row {row_number}: date does not match its voucher."
            )
        reference = None if _is_blank(raw_reference) else str(raw_reference).strip()
        memo = _text(raw_memo, "Stock Movements", row_number, "Memo")
        _validate_inventory_posting(
            source_id, posted_lines[source_id], item_accounts[sku], movement_type, value
        )

        stock = ledger._db.execute(
            "SELECT on_hand_milli, value_minor FROM inventory_items WHERE id = ?",
            (item_ids[sku],),
        ).fetchone()
        if movement_type == "receipt":
            next_quantity = stock["on_hand_milli"] + quantity
            next_value = stock["value_minor"] + value
        else:
            if quantity > stock["on_hand_milli"]:
                raise AccountingError(
                    f"Stock Movements row {row_number}: issue exceeds available stock."
                )
            expected_value = (
                stock["value_minor"]
                if quantity == stock["on_hand_milli"]
                else Ledger._round_ratio(stock["value_minor"] * quantity, stock["on_hand_milli"])
            )
            if value != expected_value:
                raise AccountingError(
                    f"Stock Movements row {row_number}: issue value does not match moving average."
                )
            if quantity < stock["on_hand_milli"] and value >= stock["value_minor"]:
                raise AccountingError(
                    f"Stock Movements row {row_number}: issue consumes all value but not all stock."
                )
            next_quantity = stock["on_hand_milli"] - quantity
            next_value = stock["value_minor"] - value

        ledger._db.execute(
            "UPDATE inventory_items SET on_hand_milli = ?, value_minor = ? WHERE id = ?",
            (next_quantity, next_value, item_ids[sku]),
        )
        ledger._db.execute(
            """INSERT INTO inventory_movements
               (item_id, entry_id, movement_date, movement_type, quantity_milli,
                value_minor, reference, memo)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                item_ids[sku], source_entries[source_id], movement_date.isoformat(),
                movement_type, quantity, value, reference, memo,
            ),
        )

    if len(movement_seen) != len(movement_rows):
        raise AccountingError("Some stock movements could not be imported.")
    for sku, (expected_quantity, expected_value) in expected.items():
        actual = ledger._db.execute(
            "SELECT on_hand_milli, value_minor FROM inventory_items WHERE sku = ?", (sku,)
        ).fetchone()
        if (actual["on_hand_milli"], actual["value_minor"]) != (
            expected_quantity, expected_value
        ):
            raise AccountingError(
                f"Inventory Items row for {sku}: exported balances do not match movement history."
            )


def _validate_inventory_posting(
    source_id: int,
    lines: list[tuple[str, int, int]],
    accounts: tuple[str, str],
    movement_type: str,
    value_minor: int,
) -> None:
    inventory_code, cogs_code = accounts
    if len(lines) != 2:
        raise AccountingError(
            f"Stock voucher {source_id} must contain exactly two journal lines."
        )
    if movement_type == "receipt":
        expected = sorted(((inventory_code, value_minor, 0),))
        actual_inventory = sorted(
            (code, debit, credit)
            for code, debit, credit in lines
            if code == inventory_code
        )
        if actual_inventory != expected or sum(line[2] for line in lines) != value_minor:
            raise AccountingError(
                f"Receipt voucher {source_id} does not debit inventory for its stock value."
            )
    else:
        if sorted(lines) != sorted(
            ((cogs_code, value_minor, 0), (inventory_code, 0, value_minor))
        ):
            raise AccountingError(
                f"Issue voucher {source_id} does not post matching inventory and COGS."
            )


def _account_id(ledger: Ledger, code: str, row_number: int, label: str) -> int:
    result = ledger._db.execute("SELECT id FROM accounts WHERE code = ?", (code,)).fetchone()
    if result is None:
        raise AccountingError(
            f"Inventory Items row {row_number}: unknown {label} account {code}."
        )
    return result["id"]


def _quantity(value: object, row_number: int, column: str) -> Decimal:
    try:
        if isinstance(value, bool) or value is None:
            raise ValueError
        return Decimal(str(value))
    except (ArithmeticError, ValueError, TypeError) as error:
        raise AccountingError(f"Inventory Items row {row_number}: invalid {column}.") from error


def _date(value: object, sheet: str, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            result = date.fromisoformat(value)
        except ValueError as error:
            raise AccountingError(
                f"{sheet} row {row_number}: date must use YYYY-MM-DD."
            ) from error
        if result.isoformat() == value:
            return result
    raise AccountingError(f"{sheet} row {row_number}: invalid date.")


def _source_id(
    value: object, sheet: str, row_number: int, label: str = "Source ID"
) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise AccountingError(f"{sheet} row {row_number}: invalid {label}.")
    try:
        identifier = int(value)
    except (ValueError, OverflowError) as error:
        raise AccountingError(f"{sheet} row {row_number}: invalid {label}.") from error
    if identifier != value or identifier < 1:
        raise AccountingError(f"{sheet} row {row_number}: invalid {label}.")
    return identifier


def _text(value: object, sheet: str, row_number: int, column: str) -> str:
    if _is_blank(value):
        raise AccountingError(f"{sheet} row {row_number}: {column} is required.")
    return str(value).strip()


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())
